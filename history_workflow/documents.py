"""Readable, versioned execution documents generated from authoritative state."""
import json
from pathlib import Path
from .core import write


def pretty(value):
    return json.dumps(value, ensure_ascii=False, indent=2)


def generate(folder: Path, state, stage, version):
    from docx import Document
    from docx.shared import Pt, Cm, RGBColor
    from docx.oxml.ns import qn
    label = {'PLAN':'计划稿','TIMED':'音频实测时轴版','HANDOFF':'冻结交接版','EXECUTED':'已执行待复审','FINAL':'归档定版'}[stage]
    title = f'{state["episode"]} 剪辑执行文档 v{version:03d}'
    sections = []
    defaults = state.get('editing_defaults', {})
    if defaults:
        sections.append(('项目默认剪辑参数', ['本版生成时的项目参数快照；逐镜头实际设置与安全检查仍以镜头表为准。', pretty(defaults)]))
    if state.get('cover_plan'):
        labels = '、'.join(spec['label'] for spec in state['cover_plan']['targets'].values())
        sections.append(('平台封面任务', [labels + '分别构图；锁稿后可与配音独立推进，导出后逐套核验像素尺寸。', pretty(state['cover_plan'])]))
    if state.get('visual_policy'):
        from .visuals import report
        sections.append(('正文图片与关键画面规划', [pretty(report(state))]))
    e = state['estimate']
    sections.append(('版本与时长', [f'本版状态：{label}。', f'锁稿版本：v{state["script"]["version"]}。锁稿输入校验：{state["lock"]["input_hash"]}', f'成片接受区间：{state["range"][0]} 至 {state["range"][1]} 秒。', f'文稿有效长度：{e["characters"]} 字。语速：{e["rate"]:.2f} 字/分钟，依据：{e["basis"]}。', f'预计旁白：{e["audio_seconds"]:.2f} 秒。预计视频：{e["video_seconds"]:.2f} 秒。估算不能替代成片实测。']))
    audio = []
    for task in state['audio_tasks']:
        asset = state['assets'].get(task['id'], {})
        audio.extend([f'{task["id"]}　对应段落 {", ".join(task["paragraphs"])}　实测时长 {asset.get("seconds", "待测")} 秒', f'已选音色：{task["voice"]}；网页设置：{pretty(task["settings"])}', f'采用音频：{asset.get("path", "待导入")}。', f'原文：{task["original"]}', f'本任务朗读文字：{task["spoken"]}'])
    sections.append(('配音顺序与文字对应', audio or ['尚未准备网页配音包。沿用已选音色；字幕使用锁稿原文，不能使用发音替换文字。']))
    plan = state['plan'] or {}
    shots = [f'执行负责人：{plan.get("responsible", "待指定")}。']
    for shot in plan.get('shots', []):
        shots.extend([f'{shot["id"]}　源图 {shot["image"]}　段落 {", ".join(shot["paragraphs"])}', f'实测起止：{shot.get("start", "待填")} 至 {shot.get("end", "待填")} 秒。锚点依据：{shot.get("anchor_evidence", "待核对")}。', f'主运动：{shot["motion"]}。目的：{shot["purpose"]}。', f'起点裁切：{shot.get("start_crop", "待填")}；终点裁切：{shot.get("end_crop", "待填")}。', f'缓动：{shot.get("easing", "由工具小样核验后填写")}。安全复核：{shot.get("safety_review", "待复核")}。', f'检查方法：{shot["check_method"]}。'])
        if shot.get('exception'):
            shots.append('静止例外及复审：' + pretty(shot['exception']))
    if not plan:
        shots.append('镜头计划待填写。未知时间和裁切保持空白，不将估算包装为实测。')
    for t in plan.get('transitions', []):
        shots.append(f'转场 {t.get("from")} → {t.get("to")}：{t.get("type", "待填")}，重叠 {t.get("overlap", "待填")} 秒。语义依据：{t.get("semantic_reason", "待填")}。重影复核：{t.get("ghosting_review", "待填")}。旁白连续性：{t.get("audio_continuity", "待填")}。')
    sections.append(('逐镜头操作与转场', shots))
    mapping=[]
    for aid,a in state['assets'].items():
        mapping.extend([f'{aid}　类型 {a["kind"]}　{"失效需重做" if a["stale"] else "当前采用"}　{"已登记人工审核" if a["review"] else "待审核"}', f'相对路径：{a["path"]}', f'尺寸：{a.get("width", "不适用")} × {a.get("height", "不适用")}；时长：{a.get("seconds", "不适用")} 秒。', f'来源：{a["source"]}；使用条件：{a["rights"]}。', f'文件校验 SHA256：{a["sha256"]}'])
    sections.append(('素材对应表', mapping or ['尚未导入素材。']))
    sources=[]
    for sid,s in state['sources'].items():
        sources.extend([f'{sid}　{s["title"]}　体裁 {s["genre"]}', f'定位：{s["locator"]}。获取状态：{s["access"]}。获取日期：{s["obtained_at"]}。', f'署名或作者：{s.get("author", "待核对")}。链接：{s.get("url", "未提供")}。'])
    sections.append(('资料来源', sources))
    sections.append(('字幕音轨与交接要求', ['按 A 编号顺序使用确认旁白；不改写、不换声、不自动删除评书停顿。', '字幕以锁稿正确文字为准。尚无真实逐字对齐；由外部工具识别或导入可靠字幕，并复查人名、年代及限定语。', '正文使用少量代表场景背景图，30分钟默认参考6张，封面另算。默认自然缓慢运动，静止必须有完整例外记录。', '画面叠化不改变旁白长度。按实测旁白检查镜头覆盖并集，不重复加减转场时长。', '配乐默认关闭。外部工具按已核验能力人工设置；本包不是剪辑软件原生工程。']))
    export=state['export']
    execution=['尚未执行，实际工具设置、工程和成片待回填。']
    if export:
        execution=[f'执行人：{export["executed_by"]}。实测视频时长：{export["seconds"]:.3f} 秒。', f'视频：{export["video"]["path"]}。', f'工程：{export.get("project_url") or (export.get("project") or {}).get("path", "待填")}。', '实际设置：' + pretty(export['settings'])]
    sections.append(('实际执行与成片', execution))
    sections.append(('最终复审', [pretty(state['accepted']) if state['accepted'] else '未完成最终复审。导出成功不等于验收通过。']))
    changes=[f'{event["time"]}　r{event["revision"]}　{event["action"]}' for event in state['events']]
    sections.append(('审核与修改记录', changes))
    doc = Document()
    section = doc.sections[0]
    section.top_margin = section.bottom_margin = Cm(2)
    section.left_margin = section.right_margin = Cm(2)
    for name in ('Normal','Title','Heading 1'):
        style=doc.styles[name]
        style.font.name='Arial'
        style._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Hiragino Sans GB')
        style.font.color.rgb=RGBColor(0,0,0)
        fonts = style._element.get_or_add_rPr().rFonts
        for key in list(fonts.attrib):
            if key.endswith('Theme') or key.endswith('theme'):
                del fonts.attrib[key]
        for border in style._element.xpath('./w:pPr/w:pBdr'):
            border.getparent().remove(border)
    doc.styles['Normal'].font.size=Pt(10)
    doc.styles['Normal'].paragraph_format.space_after=Pt(5)
    doc.add_paragraph(title,'Title')
    intro='本文件供剪辑执行人与复审人使用，记录本集采用版本、素材、镜头操作及验收状态。未取得的实测结果保持待填写；以本目录结构快照为版本依据。'
    doc.add_paragraph(intro)
    lines=[title,intro]
    for heading, paragraphs in sections:
        doc.add_heading(heading,level=1)
        lines.extend(['',heading])
        for paragraph in paragraphs:
            doc.add_paragraph(paragraph);lines.append(paragraph)
    docx=folder/f'本集剪辑执行文档_v{version:03d}.docx';doc.save(docx)
    txt=folder/f'本集剪辑执行文档_v{version:03d}.txt';write(txt,'\n'.join(lines))
    mapping_file=folder/'素材对应表.txt';write(mapping_file,'\n'.join(mapping))
    changes_file=folder/'剪辑审核与修改记录.txt';write(changes_file,'\n'.join(changes))
    manifest=folder/'结构快照.json';write(manifest,{**{k:state[k] for k in ('lock','estimate','audio_tasks','assets','plan','export','accepted','events')}, 'editing_defaults': defaults, 'cover_plan': state.get('cover_plan')})
    return [docx,txt,mapping_file,changes_file,manifest]
